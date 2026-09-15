#!/usr/bin/env python3
"""Build the architecture-group Q&A deck (HTML + PPTX) from one slide list."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import nsmap
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT_HTML = ROOT / "docs/design/arch-qa-deck.html"
OUT_PPTX = ROOT / "docs/design/arch-qa-deck.pptx"

NAVY = RGBColor(0x09, 0x0D, 0x1C)
SURFACE = RGBColor(0x16, 0x1D, 0x33)
INK = RGBColor(0xE6, 0xE8, 0xEE)
MUTED = RGBColor(0xC7, 0xCA, 0xD6)
ACCENT = RGBColor(0x81, 0x8C, 0xF8)
OK = RGBColor(0x34, 0xD3, 0x99)
WARN = RGBColor(0xFB, 0xBF, 0x24)
FAIL = RGBColor(0xF8, 0x71, 0x71)

SLIDES = [
    {
        "kicker": "Architecture group  ·  CEO / CIO room",
        "title": "One Finance UX",
        "lede": "A thin on-prem outcome layer. Subscribe to facts. Do not replace Motif, Helix, or the bus.",
        "points": [
            "We fold Ready / Blocked / Delayed for one question per outcome.",
            "Kafka, SNS, Redis, Airflow, Lambda — adjacent tools, not this product.",
            "Phase 1 runs as-is. No new cloud bill. No new broker licence.",
        ],
    },
    {
        "kicker": "The ask",
        "title": "Endorse the fold. Reuse the estate.",
        "points": [
            "Build the hub + entitled console in-house — the only Ready/Blocked system of decision.",
            "Reuse Kafka / Solace / IBM MQ if the bank already has them. Do not buy AWS to get a close.",
            "Next outcome is data. First production outcome: FOBO (Helix on Reports + rec on the Board).",
        ],
    },
    {
        "kicker": "The problem",
        "title": "The answer is scattered.",
        "lede": "“Is it safe to run this process yet?” lives across Motif, CATS, MBR, Helix, SAP.",
        "points": [
            "Each system has its own screen and its own word for done.",
            "People chase. Closes run late. Audit asks who knew — nobody holds the picture.",
            "Without this layer, every team writes a different Ready/Blocked.",
        ],
    },
    {
        "kicker": "What this is",
        "title": "Build the decision. Reuse the pipes.",
        "pairs": [
            ("We build", "Ingest a fact → distinct-key fold → command once when READY → entitle 404 → audit the human step → one board."),
            ("We reuse", "Motif, Helix, CATS, CEES, Kafka/MQ already paid, Oracle, Airflow after READY."),
            ("We never", "Rebuild books grids. Kafka in the browser. if (FOBO). LLM on the fold."),
        ],
    },
    {
        "kicker": "Category error",
        "title": "They move bytes. We decide Ready.",
        "table": [
            ("Tool they name", "Good at", "Not a close"),
            ("Kafka / IBM MQ / AMQ", "Durable log, fan-out", "A topic offset is not R-2031 BLOCKED on MB014"),
            ("SNS + SQS / Service Bus", "Queues and pub-sub", "A receipt is not Ready"),
            ("Lambda / Functions", "Glue and burst compute", "Re-implements the hub, poorly, off-prem"),
            ("Redis", "Cache", "Second truth if it owns Ready"),
            ("Airflow / Control-M", "Batch after READY", "Polling Motif recreates the late close"),
            ("Power BI / Fabric", "Yesterday’s MI", "Not a command + sign-off"),
        ],
    },
    {
        "kicker": "Q  ·  Kafka / IBM MQ / ActiveMQ",
        "title": "Use the bus. Do not confuse it with the fold.",
        "lede": "A message on motif.book.lifecycle proves produce/consume. It does not answer FOBO.",
        "points": [
            "Kafka stores 300 MASTERBOOK_READY. The hub counts distinct keys, then POSTs Helix once.",
            "IBM MQ / AMQ is the same category — a good adapter behind POST /api/events.",
            "The browser never talks to the bus. REST + SSE only.",
        ],
    },
    {
        "kicker": "Q  ·  AWS SNS, SQS, Lambda",
        "title": "Cloud messaging is not an outcome layer.",
        "points": [
            "SNS/SQS/EventBridge/Event Grid move facts. Same gap as Kafka without the fold.",
            "Lambda / Step Functions / Azure Functions = this hub rewritten as billed invocations at COB burst.",
            "Close keys, run ids, and sign-off evidence leave the bank. That is the wrong control plane.",
        ],
    },
    {
        "kicker": "Q  ·  Redis",
        "title": "We do not use Redis. We do not need it.",
        "lede": "The fold is event-sourced. State rebuilds from the event store.",
        "points": [
            "Caching Ready/Blocked in Redis creates a second truth on failover.",
            "Optional later: read-scale only. Never the system of decision.",
            "Phase 1: no cache licence. H2 now, Oracle later — already on the estate.",
        ],
    },
    {
        "kicker": "Q  ·  Airflow / Control-M",
        "title": "Batch after READY. Never instead of READY.",
        "table": [
            ("If", "Airflow", "This fold"),
            ("Same book twice", "Easy to double-count", "Distinct sourceKey"),
            ("Book 147 fails", "Generic job error", "Named blocker; RTB replay"),
            ("Old Helix run", "Can land on today", "Stale runId ignored"),
            ("Motif 20s late", "Next poke, often minutes", "Event updates now"),
            ("Next outcome", "New DAG", "POST /definitions — data"),
        ],
    },
    {
        "kicker": "Q  ·  Zero-code phase 1",
        "title": "There is no zero-code substitute.",
        "points": [
            "ServiceNow / Power Automate / n8n can land a ticket. They cannot fold distinct keys, ignore a stale run id, or 404 an unentitled row.",
            "A vendor “control tower” RFP does not stitch REV-ACC tomorrow. The POC already does.",
            "Zero new code means N shadow folds — spreadsheets, DAGs, dashboards — not a saving.",
        ],
    },
    {
        "kicker": "Q  ·  Use it as-is",
        "title": "Yes. That is the cheapest path that works.",
        "lede": "Three processes. HTTP ingest. No Kafka. No Redis. No AWS.",
        "points": [
            "Origins POST /api/events (or a 20-line adapter). Same JSON Schema in production.",
            "Kafka / Redis / Airflow are optional later, behind that API. They do not replace the product.",
            "You save the cost of a cloud account, a new broker, and a cache you would then have to distrust.",
        ],
    },
    {
        "kicker": "Later mix — optional",
        "title": "If the bank already owns a bus, plug it in.",
        "points": [
            "Motif → Kafka topic → thin adapter → POST /api/events. The fold does not change.",
            "ActionExecutor → real Helix URL. Completion is another fact with the same runId.",
            "Airflow may load a warehouse after GENERATED. It may not poll Motif to decide Helix.",
        ],
    },
    {
        "kicker": "Why on-prem",
        "title": "The close stays in the bank.",
        "pairs": [
            ("On-prem, this app", "Facts and sign-off in the DC. Reuse paid Kafka/MQ/Oracle. Three JVM processes. Next outcome is data."),
            ("AWS / Azure / Fabric", "Egress, residency, burst bill, a second IAM plane, lock-in."),
            ("Paid control-tower", "Copies Motif/Helix into a lake. Opposite of subscribe-don’t-replace."),
        ],
    },
    {
        "kicker": "Cost",
        "title": "Build once. Or pay forever in copies.",
        "points": [
            "The alternative is not zero code. It is FOBO in a spreadsheet, 15C3 in Control-M, PnL in a checklist, RTB in tickets.",
            "One SAP fact must feed two outcomes without being counted twice. Only a shared fold does that.",
            "Value: controller minutes, SLA misses, rework, audit packs — see executive-brief.md §4.",
        ],
    },
    {
        "kicker": "Generic by design",
        "title": "The next use case is a row, not a release.",
        "points": [
            "Outcome = question + feeds + SLA + on-ready. Kit = sources + destinations + embed + verbs.",
            "FOBO, 15C3, PnL, month-end — same engine. There is no if (FOBO).",
            "New command type = one ActionExecutor bean. New kit verb = data.",
        ],
    },
    {
        "kicker": "Stay thin",
        "title": "What we refuse to build",
        "points": [
            "No Kafka/Solace brokers. No Motif books grid. No Helix recon clone.",
            "No CEES replacement. No second mobile product. No LLM on Ready/Blocked.",
            "If a design needs those, it is the wrong layer.",
        ],
    },
    {
        "kicker": "Minute this",
        "title": "Six lines for the room",
        "points": [
            "Subscribe to facts. Do not replace systems of record.",
            "Reuse the bus the bank already has. Airflow only after READY.",
            "This hub is the only Ready/Blocked system of decision.",
            "Ingest stays the POC contract. Next origin is an adapter.",
            "First production outcome: FOBO. CEES fail-closed.",
            "Non-goals stay binding so the layer stays thin.",
        ],
    },
    {
        "kicker": "Evidence, not slideware",
        "title": "Point at the working product.",
        "points": [
            "OutcomeEngineTest — distinct keys, FAILED blocks, stale runId ignored.",
            "Drive FOBO / Helix → Reports GENERATED 300/300. helix-walkthrough.mp4.",
            "Onboarding → POST /api/outcomes/definitions. No new Java type.",
            "If a counter-proposal cannot show named blocker + run-id + 404 + onboard-as-data, it is not a replacement.",
        ],
    },
]


def _set_run(run, text, size=20, bold=False, color=INK, name="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def _fill(shape, rgb):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    shape.line.fill.background()


def _add_text_box(slide, l, t, w, h, text, size=20, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    _set_run(p.add_run() if p.runs else p.runs[0] if False else p.add_run(), text, size, bold, color)
    # python-pptx: empty first run — use paragraphs properly
    return box


def _write_para(tf, text, size=20, bold=False, color=INK, space_after=8, align=PP_ALIGN.LEFT, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    run = p.add_run()
    _set_run(run, text, size, bold, color)
    return p


def build_pptx() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for i, s in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(blank)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        _fill(bg, NAVY)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), prs.slide_height)
        _fill(bar, ACCENT)

        k = slide.shapes.add_textbox(Inches(0.7), Inches(0.32), Inches(11.8), Inches(0.4))
        tf = k.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        run = p.add_run()
        _set_run(run, s["kicker"].upper(), 12, True, ACCENT)

        t = slide.shapes.add_textbox(Inches(0.7), Inches(0.7), Inches(12), Inches(1.1))
        tf = t.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        _set_run(run, s["title"], 32, True, INK)

        top = Inches(1.95)
        if s.get("lede"):
            b = slide.shapes.add_textbox(Inches(0.7), Inches(1.85), Inches(12), Inches(0.7))
            tf = b.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            _set_run(run, s["lede"], 18, False, MUTED)
            top = Inches(2.55)

        if s.get("points"):
            b = slide.shapes.add_textbox(Inches(0.7), top, Inches(12), Inches(4.4))
            tf = b.text_frame
            tf.word_wrap = True
            for j, line in enumerate(s["points"]):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.level = 0
                p.space_after = Pt(14)
                run = p.add_run()
                _set_run(run, "▸  " + line, 20, False, INK)

        if s.get("pairs"):
            y = top
            for label, body in s["pairs"]:
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), y, Inches(12), Inches(1.25))
                _fill(card, SURFACE)
                tb = slide.shapes.add_textbox(Inches(0.95), y + Inches(0.12), Inches(11.5), Inches(1.05))
                tf = tb.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                run = p.add_run()
                _set_run(run, label, 13, True, OK)
                p = tf.add_paragraph()
                run = p.add_run()
                _set_run(run, body, 18, False, INK)
                y += Inches(1.4)

        if s.get("table"):
            rows = s["table"]
            cols = len(rows[0])
            table_shape = slide.shapes.add_table(len(rows), cols, Inches(0.55), top, Inches(12.2), Inches(0.48) * len(rows))
            table = table_shape.table
            widths = [Inches(3.1), Inches(4.3), Inches(4.8)] if cols == 3 else [Inches(12.2 / cols)] * cols
            for ci, w in enumerate(widths[:cols]):
                table.columns[ci].width = w
            for ri, row in enumerate(rows):
                for ci, cell_text in enumerate(row):
                    cell = table.cell(ri, ci)
                    cell.text = ""
                    p = cell.text_frame.paragraphs[0]
                    run = p.add_run()
                    header = ri == 0
                    _set_run(run, cell_text, 13 if not header else 12, header, ACCENT if header else INK)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = SURFACE if ri % 2 else RGBColor(0x12, 0x18, 0x2C)

        foot = slide.shapes.add_textbox(Inches(0.7), Inches(7.1), Inches(10), Inches(0.28))
        tf = foot.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        _set_run(run, "One Finance UX  ·  confidential  ·  architecture Q&A", 11, False, MUTED)
        num = slide.shapes.add_textbox(Inches(11.8), Inches(7.1), Inches(1.1), Inches(0.28))
        tf = num.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        run = p.add_run()
        _set_run(run, f"{i} / {len(SLIDES)}", 11, False, MUTED)

    prs.save(OUT_PPTX)
    print("wrote", OUT_PPTX)


def build_html() -> None:
    parts = []
    for i, s in enumerate(SLIDES, 1):
        inner = [f'<p class="kicker">{s["kicker"]}</p>', f"<h1>{s['title']}</h1>"]
        if s.get("lede"):
            inner.append(f'<p class="lede">{s["lede"]}</p>')
        if s.get("points"):
            inner.append("<ul>" + "".join(f"<li>{p}</li>" for p in s["points"]) + "</ul>")
        if s.get("pairs"):
            cards = "".join(
                f'<div class="card"><div class="lbl">{a}</div><p>{b}</p></div>' for a, b in s["pairs"]
            )
            inner.append(f'<div class="cards">{cards}</div>')
        if s.get("table"):
            rows = s["table"]
            head = "".join(f"<th>{c}</th>" for c in rows[0])
            body = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in rows[1:])
            inner.append(f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")
        inner.append(f'<div class="foot"><span>One Finance UX · architecture Q&A</span><span>{i} / {len(SLIDES)}</span></div>')
        parts.append(f'<section class="slide" id="s{i}">{"".join(inner)}</section>')

    html = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<title>One Finance UX — architecture Q&A</title>
<style>
  :root {{
    --canvas:#090d1c; --surface:#161d33; --ink:#e6e8ee; --muted:#c7cad6;
    --accent:#818cf8; --ok:#34d399; --stroke:#232b48;
    --sans: Inter, Segoe UI, system-ui, sans-serif;
  }}
  * {{ box-sizing:border-box; }}
  html,body {{ margin:0; height:100%; background:var(--canvas); color:var(--ink); font-family:var(--sans); }}
  .deck {{ height:100%; overflow:hidden; }}
  .slide {{
    display:none; height:100vh; padding:56px 72px 48px;
    flex-direction:column; border-left:6px solid var(--accent);
  }}
  .slide.on {{ display:flex; }}
  .kicker {{ color:var(--accent); letter-spacing:.14em; text-transform:uppercase; font-size:13px; font-weight:700; margin:0 0 10px; }}
  h1 {{ font-size:42px; line-height:1.15; margin:0 0 18px; font-weight:700; }}
  .lede {{ color:var(--muted); font-size:22px; line-height:1.4; margin:0 0 22px; max-width:52rem; }}
  ul {{ margin:8px 0 0; padding:0; list-style:none; }}
  li {{ font-size:22px; line-height:1.4; margin:0 0 16px; padding-left:28px; position:relative; max-width:54rem; }}
  li:before {{ content:"▸"; position:absolute; left:0; color:var(--accent); }}
  .cards {{ display:flex; flex-direction:column; gap:14px; margin-top:8px; }}
  .card {{ background:var(--surface); border:1px solid var(--stroke); border-radius:12px; padding:16px 20px; }}
  .card .lbl {{ color:var(--ok); font-size:13px; font-weight:700; letter-spacing:.08em; text-transform:uppercase; margin-bottom:6px; }}
  .card p {{ margin:0; font-size:20px; line-height:1.4; }}
  table {{ border-collapse:collapse; width:100%; margin-top:8px; font-size:16px; }}
  th,td {{ text-align:left; padding:10px 12px; border-bottom:1px solid var(--stroke); vertical-align:top; }}
  th {{ color:var(--accent); font-size:12px; letter-spacing:.08em; text-transform:uppercase; }}
  .foot {{ margin-top:auto; display:flex; justify-content:space-between; color:var(--muted); font-size:13px; padding-top:18px; }}
  .hint {{ position:fixed; right:28px; bottom:22px; color:#7e85a3; font-size:12px; }}
  @media print {{
    .slide {{ display:flex; height:100vh; page-break-after:always; }}
    .hint {{ display:none; }}
  }}
</style>
</head>
<body>
<div class="deck">
{"".join(parts)}
</div>
<div class="hint">← →  or click  ·  F for fullscreen</div>
<script>
const slides=[...document.querySelectorAll('.slide')];
let i=0;
function show(n){{ i=(n+slides.length)%slides.length; slides.forEach((s,k)=>s.classList.toggle('on',k===i)); history.replaceState(null,'','#s'+(i+1)); }}
document.addEventListener('keydown',e=>{{
  if(['ArrowRight','PageDown',' ','Enter'].includes(e.key)) {{ e.preventDefault(); show(i+1); }}
  if(['ArrowLeft','PageUp','Backspace'].includes(e.key)) {{ e.preventDefault(); show(i-1); }}
  if(e.key==='Home') show(0);
  if(e.key==='End') show(slides.length-1);
  if(e.key==='f'||e.key==='F') document.documentElement.requestFullscreen?.();
}});
document.addEventListener('click',e=>{{ if(!e.target.closest('a')) show(i+1); }});
const start=parseInt((location.hash||'#s1').slice(2),10)-1;
show(Number.isFinite(start)&&start>=0?start:0);
</script>
</body>
</html>
"""
    OUT_HTML.write_text(html)
    print("wrote", OUT_HTML)


if __name__ == "__main__":
    build_html()
    build_pptx()
